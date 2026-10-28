import csv
import io
import shutil
import tempfile
from pathlib import Path

from django.test import override_settings
from django.urls import reverse

from core.tests import BaseTestCase

from .datasets import all_datasets
from .models import ExportJob

TMP = Path(tempfile.mkdtemp())


@override_settings(EXPORT_RUNNER='inline', EXPORT_ROOT=TMP)
class ExportTests(BaseTestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TMP, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.patient = self.make_patient(first_name='Zuhura')
        self.manager = self.make_user('dm_x', 'Data Manager', facility=self.ilala)
        self.pi = self.make_user('pi_x', 'PI / Co-Investigator', level='national')

    def export(self, user, **data):
        self.client.force_login(user)
        payload = {'dataset': 'patients', 'fmt': 'csv', 'mode': 'labels'}
        payload.update(data)
        self.client.post(reverse('exports:list'), payload)
        return ExportJob.objects.filter(user=user).latest('pk')

    def read_csv(self, job):
        return list(csv.reader(io.StringIO(job.path.read_text(encoding='utf-8-sig'))))

    def test_pi_export_is_deidentified(self):
        job = self.export(self.pi)
        self.assertEqual((job.status, job.deidentified, job.rows), ('done', True, 1))
        header, row = self.read_csv(job)
        self.assertIn('patient_id', header)
        self.assertNotIn('first_name', header)
        self.assertNotIn('Zuhura', row)
        self.assertEqual(self.client.get(reverse('exports:download', args=[job.pk])).status_code, 200)

    def test_data_manager_identifiable_only_when_ticked_and_download_is_private(self):
        job = self.export(self.manager)
        self.assertNotIn('first_name', self.read_csv(job)[0])
        job = self.export(self.manager, identifiable='on')
        header, row = self.read_csv(job)
        self.assertIn('Zuhura', row)
        self.assertFalse(job.deidentified)
        self.client.force_login(self.pi)
        self.assertEqual(self.client.get(reverse('exports:download', args=[job.pk])).status_code, 404)

    def test_scope_limits_rows_and_codes_mode(self):
        self.make_patient(self.mwanza, first_name='Other')
        job = self.export(self.manager, mode='codes')
        rows = self.read_csv(job)
        self.assertEqual(len(rows), 2)               # header + the one Ilala patient
        self.assertIn('female', rows[1])

    def test_all_datasets_workbook_and_every_dataset_builds(self):
        from openpyxl import load_workbook
        job = self.export(self.admin, dataset='all', fmt='csv')
        self.assertEqual((job.status, job.fmt), ('done', 'xlsx'), job.error)
        workbook = load_workbook(job.path, read_only=True)
        self.assertIn('Codebook', workbook.sheetnames)
        self.assertGreaterEqual(len(workbook.sheetnames), len(all_datasets()))

    def test_codebook_formats(self):
        self.client.force_login(self.pi)
        self.assertContains(self.client.get(reverse('exports:codebook')), 'patient_id')
        for fmt, ctype in (('csv', 'text/csv'), ('xlsx', 'spreadsheet'), ('pdf', 'application/pdf')):
            with self.subTest(fmt=fmt):
                response = self.client.get(reverse('exports:codebook') + f'?dataset=patients&format={fmt}')
                self.assertIn(ctype, response['Content-Type'])
        body = self.client.get(reverse('exports:codebook') + '?dataset=patients&format=csv').content.decode()
        self.assertNotIn('first_name', body)          # de-identified codebook for the PI

    def test_cleanup_removes_old_files(self):
        from django.utils import timezone

        from .tasks import cleanup_exports
        job = self.export(self.pi)
        ExportJob.objects.filter(pk=job.pk).update(finished_at=timezone.now() - timezone.timedelta(days=30))
        self.assertEqual(cleanup_exports(), 1)
        self.assertFalse(job.path.exists())

    def test_nurse_cannot_export(self):
        nurse = self.make_user('nurse_x', 'Nurse / Triage Officer', facility=self.ilala)
        self.client.force_login(nurse)
        self.assertEqual(self.client.get(reverse('exports:list')).status_code, 403)


@override_settings(EXPORT_RUNNER='inline', EXPORT_ROOT=TMP)
class FormsPageTests(BaseTestCase):
    @classmethod
    def setUpTestData(cls):
        from django.core.management import call_command
        super().setUpTestData()
        call_command('setup_penplus', stdout=open('/dev/null', 'w'))

    def setUp(self):
        from django.utils import timezone

        from laboratory.models import LabOrder, LabTest
        self.patient = self.make_patient(first_name='Zuhura')
        self.pi = self.make_user('pi_f', 'PI / Co-Investigator', level='national')
        when = timezone.now().replace(hour=9)
        for code, raw in (('hb', '9.1'), ('creat', '80'), ('hiv', 'Negative')):
            order = LabOrder(patient=self.patient, facility=self.ilala, test=LabTest.objects.get(code=code),
                             ordered_by=self.admin, ordered_at=when)
            order.record_result(raw, self.admin, when=when)
            order.save()

    def download(self, user, **data):
        self.client.force_login(user)
        self.client.post(reverse('exports:form_download'), data)
        return ExportJob.objects.filter(user=user).latest('pk')

    def test_forms_page_lists_forms_with_records(self):
        self.client.force_login(self.pi)
        response = self.client.get(reverse('exports:forms'))
        self.assertContains(response, 'Laboratory')
        self.assertContains(response, 'Radiology / imaging')
        self.assertContains(response, 'Treatment plan &amp; medicines')
        self.assertContains(response, 'lab_details')

    def test_laboratory_form_has_one_row_per_lab_date(self):
        from openpyxl import load_workbook
        job = self.download(self.pi, form='laboratory')
        self.assertEqual((job.status, job.fmt, job.deidentified), ('done', 'xlsx', True), job.error)
        sheets = load_workbook(job.path, read_only=True)
        rows = list(sheets[sheets.sheetnames[0]].values)
        self.assertEqual(len(rows), 2)                       # header + one lab date with three tests
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual((row['hb'], row['creat'], row['hiv']), (9.1, 80, 'Negative'))
        self.assertEqual(row['patient_id'], self.patient.patient_number)
        self.assertEqual(len(rows[0]), len(set(rows[0])))    # no duplicate column names

    def test_one_table_as_csv_and_whole_study(self):
        job = self.download(self.pi, form='laboratory', dataset='labs')
        self.assertEqual((job.status, job.fmt, job.dataset), ('done', 'csv', 'labs'), job.error)
        job = self.download(self.pi, form='*')
        self.assertEqual((job.status, job.fmt), ('done', 'xlsx'), job.error)
        self.assertNotIn(':', job.file_name)

    def test_form_dictionary(self):
        self.client.force_login(self.pi)
        self.assertContains(self.client.get(reverse('exports:codebook') + '?dataset=form:laboratory'), 'hb')
        for fmt in ('xlsx', 'pdf'):
            response = self.client.get(reverse('exports:codebook') + f'?dataset=form:treatment&format={fmt}')
            self.assertEqual(response.status_code, 200)
            self.assertIn('attachment', response['Content-Disposition'])

    def test_disease_limited_staff_export_only_their_programme(self):
        from clinical.models import ClinicalProgramme, ProgrammeEnrollment
        ProgrammeEnrollment.objects.create(patient=self.patient, facility=self.ilala,
                                           programme=ClinicalProgramme.objects.get(code='dm-care'))
        cardiologist = self.make_user('cardio_f', 'PI / Co-Investigator', level='national')
        cardiologist.diseases.add(self.cardiac)
        job = self.download(cardiologist, form='laboratory', dataset='labs')
        self.assertEqual((job.status, job.rows), ('done', 0), job.error)
        job = self.download(self.pi, form='laboratory', dataset='labs')
        self.assertEqual(job.rows, 3)
