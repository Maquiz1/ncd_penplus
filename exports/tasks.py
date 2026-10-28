import csv
import re
import threading

from celery import shared_task
from django.conf import settings
from django.db import close_old_connections
from django.utils import timezone

from . import datasets as ds
from .models import ExportJob

PROGRESS_EVERY = 500


def file_name(job):
    tag = '_deidentified' if job.deidentified else ''
    name = job.dataset.replace('form:*', 'all_forms').replace('form:', 'form_')
    return f'NCD_{name}_{job.mode}{tag}_{timezone.localdate():%Y%m%d}_u{job.user_id}_j{job.pk}.{job.fmt}'


def selected(job):
    if job.dataset == 'all':
        return ds.all_datasets()
    if job.dataset.startswith('form:'):
        from .study_forms import form_datasets
        return form_datasets(job.dataset[5:])
    return [ds.get_dataset(job.dataset)]


def job_facilities(job):
    from facilities.scope import facilities_for_user
    facilities = facilities_for_user(job.user)       # re-checked when the file is built
    ids = job.filters.get('facilities')
    return facilities.filter(pk__in=ids) if ids else facilities


def rows_for(job, dataset, queryset):
    columns = ds.visible_columns(dataset, job.deidentified)
    for obj in queryset.iterator(chunk_size=1000):
        yield [ds.plain(c.getter(obj, job.mode)) for c in columns]


def sheet_title(label, used):
    title = re.sub(r'[\[\]:*?/\\]', '', label)[:28] or 'Sheet'
    base, n = title, 2
    while title in used:
        title = f'{base[:25]}_{n}'
        n += 1
    used.add(title)
    return title


def build(job):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    from facilities.scope import disease_ids_for_user
    facilities = job_facilities(job)
    start, end = job.filters.get('start'), job.filters.get('end')
    diseases = disease_ids_for_user(job.user)
    parts = [(d, ds.dataset_queryset(d, facilities, start, end, diseases)) for d in selected(job) if d is not None]
    job.total = sum(qs.count() for _d, qs in parts)
    job.status, job.done = ExportJob.Status.RUNNING, 0
    job.save(update_fields=['total', 'status', 'done'])
    settings.EXPORT_ROOT.mkdir(parents=True, exist_ok=True)
    job.file_name = file_name(job)
    path = settings.EXPORT_ROOT / job.file_name

    def tick():
        job.done += 1
        if job.done % PROGRESS_EVERY == 0:
            ExportJob.objects.filter(pk=job.pk).update(done=job.done)

    if job.fmt == ExportJob.Format.CSV:
        dataset, queryset = parts[0]
        with open(path, 'w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.writer(handle)
            writer.writerow([c.name for c in ds.visible_columns(dataset, job.deidentified)])
            for row in rows_for(job, dataset, queryset):
                writer.writerow(row)
                tick()
    else:
        workbook = Workbook(write_only=True)
        bold, used = Font(bold=True), set()
        for dataset, queryset in parts:
            sheet = workbook.create_sheet(sheet_title(dataset.label, used))
            header = []
            for c in ds.visible_columns(dataset, job.deidentified):
                cell = WriteOnlyCell(sheet, value=c.name)
                cell.font = bold
                header.append(cell)
            sheet.append(header)
            for row in rows_for(job, dataset, queryset):
                sheet.append(row)
                tick()
        codebook = workbook.create_sheet('Codebook')
        codebook.append(['Dataset', 'Variable', 'Label', 'Type', 'Unit', 'Code', 'Code label', 'In de-identified export'])
        for row in ds.codebook_rows([d for d, _q in parts], job.deidentified):
            codebook.append(row)
        workbook.save(path)
    job.rows = job.done
    job.file_size = path.stat().st_size
    job.status, job.finished_at = ExportJob.Status.DONE, timezone.now()
    job.save()


@shared_task(name='exports.build_export')
def build_export(job_id):
    job = ExportJob.objects.select_related('user').get(pk=job_id)
    try:
        build(job)
    except Exception as exc:  # noqa: BLE001 - shown to the user on the exports page
        ExportJob.objects.filter(pk=job.pk).update(status=ExportJob.Status.FAILED, error=str(exc)[:2000],
                                                   finished_at=timezone.now())
        raise


def start(job):
    """Run the export with the configured runner (Celery worker, a background thread, or inline)."""
    runner = settings.EXPORT_RUNNER
    if runner == 'inline':
        try:
            build_export(job.pk)
        except Exception:  # noqa: BLE001 - recorded on the job
            pass
    elif runner == 'thread':
        def work():
            close_old_connections()
            try:
                build_export(job.pk)
            except Exception:  # noqa: BLE001 - recorded on the job
                pass
            finally:
                close_old_connections()
        threading.Thread(target=work, daemon=True).start()
    else:
        build_export.delay(job.pk)


@shared_task(name='exports.cleanup_exports')
def cleanup_exports():
    """Remove export files older than EXPORT_RETENTION_DAYS."""
    cutoff = timezone.now() - timezone.timedelta(days=settings.EXPORT_RETENTION_DAYS)
    removed = 0
    for job in ExportJob.objects.filter(status=ExportJob.Status.DONE, finished_at__lt=cutoff):
        if job.path and job.path.exists():
            job.path.unlink()
        job.status = ExportJob.Status.EXPIRED
        job.save(update_fields=['status'])
        removed += 1
    return removed
