import csv
import io

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views import View

from audit.models import AuditEvent
from audit.services import log_event
from facilities.scope import scope_description

from . import datasets as ds
from . import study_forms
from .forms import ExportForm
from .models import ExportJob
from .tasks import start


def can_identify(user):
    return user.has_perm('exports.export_identifiable') and user.has_perm('patients.view_patient_identity')


class ExportListView(PermissionRequiredMixin, View):
    permission_required = 'reporting.export_data'

    def get(self, request, form=None):
        jobs = ExportJob.objects.filter(user=request.user)[:30]
        return render(request, 'exports/export_list.html', {
            'form': form or ExportForm(user=request.user), 'jobs': jobs, 'scope': scope_description(request.user),
            'retention': settings.EXPORT_RETENTION_DAYS,
            'deidentified_only': not can_identify(request.user), 'running': any(j.status in ('queued', 'running') for j in jobs)})

    def post(self, request):
        form = ExportForm(request.POST, user=request.user)
        if not form.is_valid():
            return self.get(request, form)
        data = form.cleaned_data
        dataset = None if data['dataset'] == 'all' else ds.get_dataset(data['dataset'])
        if data['dataset'] != 'all' and dataset is None:
            raise Http404
        facilities = form.facilities(request.user)
        label_parts = [str(data.get('site') or data.get('district') or 'All my sites')]
        if data.get('start') or data.get('end'):
            label_parts.append(f'{data.get("start") or "…"} – {data.get("end") or "…"}')
        job = ExportJob.objects.create(
            user=request.user, dataset=data['dataset'], dataset_label=dataset.label if dataset else 'All datasets',
            fmt=data['fmt'], mode=data['mode'], deidentified=not data.get('identifiable'),
            filters={'facilities': list(facilities.values_list('pk', flat=True)) if facilities is not None else None,
                     'start': data['start'].isoformat() if data.get('start') else None,
                     'end': data['end'].isoformat() if data.get('end') else None},
            filters_label=' · '.join(label_parts))
        log_event(request=request, action=AuditEvent.Action.EXPORT, obj=job,
                  changes={'dataset': job.dataset, 'deidentified': job.deidentified, 'filters': job.filters_label})
        start(job)
        messages.success(request, f'Export started: {job.dataset_label}. It appears below when ready.')
        return redirect('exports:list')


class ExportStatusView(PermissionRequiredMixin, View):
    permission_required = 'reporting.export_data'

    def get(self, request):
        jobs = ExportJob.objects.filter(user=request.user, pk__in=request.GET.getlist('id'))
        return JsonResponse({'jobs': [{'id': j.pk, 'status': j.status, 'label': j.get_status_display(),
                                       'progress': j.progress, 'done': j.done, 'total': j.total, 'error': j.error[:300]}
                                      for j in jobs]})


class ExportDownloadView(PermissionRequiredMixin, View):
    """Only the person who made the export can download it."""

    permission_required = 'reporting.export_data'

    def get(self, request, pk):
        job = get_object_or_404(ExportJob, pk=pk, user=request.user, status=ExportJob.Status.DONE)
        if not job.path or not job.path.exists():
            raise Http404('The file is no longer available.')
        ExportJob.objects.filter(pk=job.pk).update(downloads=job.downloads + 1)
        log_event(request=request, action=AuditEvent.Action.EXPORT, obj=job, reason='Export file downloaded')
        return FileResponse(open(job.path, 'rb'), as_attachment=True, filename=job.file_name.split('_u')[0] + f'.{job.fmt}')


class ExportDeleteView(PermissionRequiredMixin, View):
    permission_required = 'reporting.export_data'

    def post(self, request, pk):
        job = get_object_or_404(ExportJob, pk=pk, user=request.user)
        if job.path and job.path.exists():
            job.path.unlink()
        job.status = ExportJob.Status.EXPIRED
        job.save(update_fields=['status'])
        messages.success(request, 'Export file removed.')
        return redirect('exports:list')


class CodebookView(PermissionRequiredMixin, View):
    """Data dictionary: every variable, its label, type, unit and codes – matching the exports."""

    permission_required = 'reporting.export_data'

    def get(self, request):
        datasets = ds.all_datasets()
        key = request.GET.get('dataset', '')
        if key.startswith('form:'):
            selected = [d for d in study_forms.form_datasets(key[5:]) if d is not None]
        else:
            selected = datasets if key in ('', 'all') else [d for d in datasets if d.key == key]
        deidentified = not can_identify(request.user)
        fmt = request.GET.get('format')
        header = ['Dataset', 'Variable', 'Label', 'Type', 'Unit', 'Code', 'Code label', 'In de-identified export']
        if fmt in ('csv', 'xlsx', 'pdf'):
            rows = list(ds.codebook_rows(selected, deidentified))
            name = f'NCD_codebook_{(key or "all").replace("form:", "form_").replace("*", "all_forms")}_{timezone.localdate():%Y%m%d}'
            log_event(request=request, action=AuditEvent.Action.EXPORT, model='exports.Codebook',
                      object_repr=f'Codebook {key or "all"} ({fmt})')
            if fmt == 'csv':
                buffer = io.StringIO()
                writer = csv.writer(buffer)
                writer.writerow(header)
                writer.writerows(rows)
                response = HttpResponse('﻿' + buffer.getvalue(), content_type='text/csv; charset=utf-8')
            elif fmt == 'xlsx':
                from openpyxl import Workbook
                from openpyxl.styles import Font
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = 'Codebook'
                sheet.append(header)
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                for row in rows:
                    sheet.append(row)
                sheet.freeze_panes = 'A2'
                buffer = io.BytesIO()
                workbook.save(buffer)
                response = HttpResponse(buffer.getvalue(),
                                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            else:
                from reporting.pdf import pdf_response
                return pdf_response(request, 'exports/_codebook_body.html',
                                    {'datasets': selected, 'deidentified': deidentified}, f'Codebook {key or "all datasets"}')
            response['Content-Disposition'] = f'attachment; filename="{name}.{fmt}"'
            return response
        return render(request, 'exports/codebook.html', {
            'datasets': datasets, 'selected': selected, 'key': key, 'deidentified': deidentified,
            'body_template': 'exports/_codebook_body.html'})


class FormsView(PermissionRequiredMixin, View):
    """Every study form (as in the old EDCS) with its records in the user's scope, data downloads and dictionary."""

    permission_required = 'reporting.export_data'

    def get(self, request):
        from facilities.scope import disease_ids_for_user, facilities_for_user
        facilities, diseases = facilities_for_user(request.user), disease_ids_for_user(request.user)
        forms = study_forms.study_forms()
        for form in forms:
            form.tables = [{'dataset': d, 'records': ds.dataset_queryset(d, facilities, disease_ids=diseases).count(),
                            'variables': len(ds.visible_columns(d, True))} for d in form.datasets]
            form.records = form.tables[0]['records']
        return render(request, 'exports/forms.html', {
            'forms': forms, 'scope': scope_description(request.user), 'diseases': diseases,
            'can_identify': can_identify(request.user)})


class FormExportView(PermissionRequiredMixin, View):
    """Start a de-identified export of one form (Excel: all its tables) or one of its tables (CSV)."""

    permission_required = 'reporting.export_data'

    def post(self, request):
        key, fmt = request.POST.get('form', ''), request.POST.get('fmt', 'xlsx')
        mode = request.POST.get('mode', 'labels')
        dataset_key = request.POST.get('dataset', '')
        if fmt not in ExportJob.Format.values or mode not in ExportJob.Mode.values:
            raise Http404
        if dataset_key:
            form = study_forms.get_form(key)
            dataset = next((d for d in (form.datasets if form else []) if d.key == dataset_key), None)
            if dataset is None:
                raise Http404
            job_key, label, fmt = dataset.key, f'{form.title}: {dataset.label}', 'csv'
        elif key == '*':
            job_key, label, fmt = 'form:*', 'All study forms', 'xlsx'
        else:
            form = study_forms.get_form(key)
            if form is None:
                raise Http404
            job_key, label = f'form:{form.key}', f'{form.title} (form)'
            if len(form.datasets) > 1:
                fmt = 'xlsx'  # one sheet per table
        job = ExportJob.objects.create(user=request.user, dataset=job_key, dataset_label=label[:200], fmt=fmt,
                                       mode=mode, deidentified=True, filters={'facilities': None},
                                       filters_label='All my sites')
        log_event(request=request, action=AuditEvent.Action.EXPORT, obj=job,
                  changes={'dataset': job.dataset, 'deidentified': True, 'filters': job.filters_label})
        start(job)
        messages.success(request, f'Export started: {label}. Download it below when it is ready.')
        return redirect('exports:list')
